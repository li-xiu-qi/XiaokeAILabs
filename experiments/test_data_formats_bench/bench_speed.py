"""
xlsx vs sqlite 查询速度实测
=======================

同一份数据，三种形态做五类等价查询的耗时对比：

  1. sqlite 无索引
  2. sqlite 带索引
  3. pandas 读入 xlsx 后在内存中查询

结论速览：索引只在点查（高选择性）上有数量级优势，
在范围过滤和分组聚合上可能反而更慢（回表随机 IO）。
sqlite 相对 xlsx 的真正优势来自「一次导入、反复查询」，
不在索引本身。

运行：
    python bench_speed.py            # 全量，约 3 分钟
    python bench_speed.py --rows 100000   # 缩小规模

依赖：
    openpyxl
    pandas
"""

import argparse
import os
import random
import sqlite3
import string
import time

from openpyxl import Workbook, load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "output")

COLS = ["id", "cat", "val1", "val2", "name", "code", "amt", "flag"]


def human(size):
    for unit in ("B", "K", "M", "G"):
        if size < 1024 or unit == "G":
            return "%.2f%s" % (size, unit)
        size /= 1024.0


def generate(path, rows):
    """生成测试数据并写成 xlsx。"""
    random.seed(31)
    name_pool = [
        "".join(random.choices(string.ascii_letters, k=10)) for _ in range(5000)
    ]
    wb = Workbook(write_only=True)
    ws = wb.create_sheet()
    ws.append(COLS)
    for i in range(rows):
        ws.append([
            i,
            "cat%d" % (i % 50),
            round(random.random() * 1000, 2),
            i * 2,
            random.choice(name_pool),
            "C%07d" % i,
            round(random.random() * 100, 2),
            random.choice(["Y", "N"]),
        ])
    wb.save(path)
    return path


def import_sqlite(db_path, xlsx_path, index_cols):
    """从 xlsx 读回数据写入 sqlite。"""
    if os.path.exists(db_path):
        os.remove(db_path)

    t0 = time.time()
    con = sqlite3.connect(db_path)
    cur = con.cursor()
    cur.execute("CREATE TABLE d (%s)" % ", ".join(
        "%s %s" % (c, "INTEGER" if c in ("id", "val2") else
                   "REAL" if c in ("val1", "amt") else "TEXT")
        for c in COLS
    ))
    wb = load_workbook(xlsx_path, read_only=True)
    ws = wb.active
    it = ws.iter_rows(values_only=True)
    next(it)
    cur.executemany("INSERT INTO d VALUES (%s)" % ",".join("?" * len(COLS)), it)
    wb.close()
    for c in index_cols:
        cur.execute("CREATE INDEX ix_%s ON d(%s)" % (c, c))
    con.commit()
    con.close()
    return time.time() - t0


def timeit(label, fn, repeat=3):
    """重复执行取平均耗时（毫秒）。"""
    t0 = time.time()
    for _ in range(repeat):
        fn()
    return (time.time() - t0) / repeat * 1000.0


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--rows", type=int, default=500000)
    parser.add_argument("--skip-pandas", action="store_true",
                        help="跳过 pandas 侧（依赖较重时用）")
    args = parser.parse_args()

    os.makedirs(OUT, exist_ok=True)
    xlsx = generate(os.path.join(OUT, "speed_data.xlsx"), args.rows)
    print("数据：%d 行 x %d 列，xlsx = %s\n" % (
        args.rows, len(COLS), human(os.path.getsize(xlsx))
    ))

    # ---- sqlite 两侧 ----
    results = {}
    for tag, idx in (("无索引", ()), ("有索引", ("id", "cat", "val1"))):
        db = os.path.join(OUT, "speed_%s.sqlite" % tag)
        cost = import_sqlite(db, xlsx, idx)
        print("导入 sqlite（%s）耗时 %.1f 秒" % (tag, cost))

        con = sqlite3.connect(db)
        cur = con.cursor()
        results[tag] = {
            "点查 id": timeit("p", lambda: cur.execute(
                "SELECT * FROM d WHERE id=?" , (args.rows // 2,)
            ).fetchall(), 1),
            "等值过滤 cat": timeit("p", lambda: cur.execute(
                "SELECT COUNT(*), AVG(amt) FROM d WHERE cat='cat7'"
            ).fetchall(), 1),
            "范围过滤 val1": timeit("p", lambda: cur.execute(
                "SELECT COUNT(*), AVG(val2) FROM d WHERE val1 BETWEEN 100 AND 200"
            ).fetchall(), 1),
            "分组聚合": timeit("p", lambda: cur.execute(
                "SELECT cat, COUNT(*), SUM(amt) FROM d GROUP BY cat"
            ).fetchall(), 1),
            "全表求和": timeit("p", lambda: cur.execute(
                "SELECT SUM(val2) FROM d"
            ).fetchall(), 1),
        }
        con.close()
        print("")

    # ---- pandas 侧 ----
    if not args.skip_pandas:
        t0 = time.time()
        import pandas as pd
        df = pd.read_excel(xlsx, engine="openpyxl")
        read_cost = time.time() - t0
        print("读入 xlsx 耗时 %.1f 秒（内存 %.0fMB）\n" % (
            read_cost, df.memory_usage(deep=True).sum() / 1048576.0
        ))
        results["pandas读xlsx"] = {
            "点查 id": timeit("p", lambda: df[df["id"] == args.rows // 2]),
            "等值过滤 cat": timeit("p", lambda: df[df["cat"] == "cat7"]
                                   .agg({"amt": ["count", "mean"]})),
            "范围过滤 val1": timeit("p", lambda: df[
                (df["val1"] >= 100) & (df["val1"] <= 200)
            ].agg({"val2": ["count", "mean"]})),
            "分组聚合": timeit("p", lambda: df.groupby("cat")["amt"]
                               .agg(["count", "sum"])),
            "全表求和": timeit("p", lambda: df["val2"].sum()),
        }

    # ---- 汇总 ----
    tags = list(results.keys())
    header = "%-16s" % "查询" + "".join("%16s" % t for t in tags)
    print("\n" + header)
    print("-" * len(header))
    for op in results[tags[0]]:
        print("%-16s" % op + "".join(
            "%15.1fms" % results[t][op] for t in tags
        ))


if __name__ == "__main__":
    main()
